"""
الهدف:
استخراج نص خام من محتوى ملف وثيقة مرفوعة (PDF/Word أساسًا) — الخطوة
الأولى اللازمة قبل أي تلخيص أو بحث دلالي أو شات بوت يعتمد على محتوى
الوثيقة (راجعي app/services/document_embedding_service.py للمستهلك
الوحيد لهذه الدالة حاليًا).

المسؤولية:
دالة واحدة عامة extract_text() تحدد طريقة الاستخراج حسب mime_type ولا
ترمي أي استثناء أبدًا — ملف تالف أو نوع غير مدعوم يرجع None بدل تعطيل
كامل عملية الرفع/توليد embedding (نفس فلسفة GeminiError بـgemini_client.py:
فشل جزئي مسجَّل، وليس استثناء يتسرب للمستخدم مباشرة).

أنواع مدعومة حاليًا:
- PDF (application/pdf) عبر pypdf — نص فقط، بدون OCR للصفحات الممسوحة
  ضوئيًا كصورة (خارج نطاق هذه المرحلة).
- Word الحديث (.docx، application/vnd.openxmlformats-officedocument.
  wordprocessingml.document) عبر python-docx — فقرات + جداول.
- نص عادي (text/plain, text/csv, إلخ. أي mime_type يبدأ بـ"text/") —
  decode مباشر.
أي نوع آخر (صور، PowerPoint، Word القديم .doc...) يرجع None حاليًا —
البحث/التلخيص لهذي الوثائق يبقى معطَّلًا لحد ما تُضاف مكتبة استخراج
مناسبة له لاحقًا، دون ما يمنع رفع الوثيقة نفسها أو استخدامها عاديًا.
"""

import io

from pypdf import PdfReader
from pypdf.errors import PdfReadError
from docx import Document as DocxDocument
from docx.opc.exceptions import PackageNotFoundError

# حد أقصى لطول النص المستخرَج (أحرف) — وثيقة ضخمة جدًا (مئات الصفحات)
# تُقتطع بدل ما تُرسَل كاملة لاحقًا لـGemini (تكلفة/حد أقصى لطلب
# الـembedding والشات). القطع من النهاية وليس الرفض الكامل — الجزء الأول
# غالبًا يحمل أهم المحتوى (عنوان/مقدمة/ملخص تنفيذي).
MAX_EXTRACTED_CHARS = 100_000

_PDF_MIME = "application/pdf"
_DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _extract_pdf(content: bytes) -> str | None:
    try:
        reader = PdfReader(io.BytesIO(content))
        pages = [page.extract_text() or "" for page in reader.pages]
    except (PdfReadError, Exception):  # noqa: BLE001 — ملف PDF تالف بأي شكل، لا نفشل الرفع بسببه
        return None
    text = "\n\n".join(p for p in pages if p.strip())
    return text or None


def _extract_docx(content: bytes) -> str | None:
    try:
        doc = DocxDocument(io.BytesIO(content))
    except (PackageNotFoundError, Exception):  # noqa: BLE001 — ملف Word تالف/غير صالح
        return None
    parts: list[str] = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n".join(parts)
    return text or None


def _extract_plain_text(content: bytes) -> str | None:
    try:
        text = content.decode("utf-8")
    except UnicodeDecodeError:
        try:
            text = content.decode("utf-8", errors="ignore")
        except Exception:  # noqa: BLE001
            return None
    return text or None


def extract_text(content: bytes, mime_type: str) -> str | None:
    """يرجع النص المستخرَج (مقتطعًا عند MAX_EXTRACTED_CHARS)، أو None لو
    النوع غير مدعوم أو الاستخراج فشل — لا يرمي أي استثناء أبدًا."""
    mime_type = (mime_type or "").lower().strip()

    if mime_type == _PDF_MIME:
        text = _extract_pdf(content)
    elif mime_type == _DOCX_MIME:
        text = _extract_docx(content)
    elif mime_type.startswith("text/"):
        text = _extract_plain_text(content)
    else:
        text = None

    if text is None:
        return None
    return text[:MAX_EXTRACTED_CHARS]
