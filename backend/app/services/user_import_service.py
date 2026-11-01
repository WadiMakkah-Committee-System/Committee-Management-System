"""
الهدف:
رفع ملف المستخدمين (users.import — طلب لاما 2026-11-01): إنشاء عدة مستخدمين
دفعة واحدة من ملف Excel بنفس أعمدة "نموذج_بيانات_المستخدمين.xlsx"، مع توليد
اسم مستخدم وكلمة مرور مؤقتة لكل واحد (تُرسَل لبريده من طبقة الـAPI).

المسؤولية:
- parse_workbook(): قراءة الملف والتحقق من بنيته (الأعمدة، الحد الأقصى للصفوف).
- import_users(): التحقق من كل الصفوف أولًا، ثم الإنشاء بمعاملة واحدة.
- build_template(): توليد نموذج فارغ بقوائم منسدلة من بيانات النظام الحالية.

قواعد العمل (قرارات لاما 2026-11-01):
- الكل أو لا شيء: أي خطأ بأي صف ← لا يُنشأ أحد، وتُرجَع كل الأخطاء برقم الصف.
- الحقول الإلزامية: الاسم الأول/الأوسط/العائلة + البريد. الإدارة والمسمى
  والدور اختيارية.
- الدور: أسماء النموذج الأربعة تُقابَل بأدوار النظام (ROLE_ALIASES)، وأي
  اسم/وصف دور نظامي آخر يُقبل كما هو. الفارغ = بلا دور. دور المسؤول الأعلى
  وأدوار اللجان لا تُسند عبر الملف أبدًا.
- المسمى الوظيفي: الموجود يُستخدم، وغير الموجود يُضاف تلقائيًا لقائمة
  المسميات (طلب لاما 2026-11-01) — لمن يملك نطاق 'all' أو صلاحية
  job_titles.create فقط، وإلا يُرفض الصف.
- الإدارة: الموجودة يُضاف لها المستخدم؛ غير الموجودة تُنشأ تلقائيًا برمز
  DEP-001... ومديرها صاحب الصف الذي دوره "مدير إدارة (أدمن)" بنفس الإدارة
  (إلزامي للإدارة الجديدة — نفس قاعدة create_department: الإدارة لها مدير).
  إدارة موجودة لها مدير + صف مدير جديد لها ← خطأ (التغيير من صفحة الإدارات).
  أكثر من مدير لنفس الإدارة بالملف ← خطأ.
- الحد الأقصى MAX_ROWS صفًا (حد Brevo المجاني 300 بريد/يوم).

الصلاحيات (مفروضة هنا لا بالواجهة):
- نطاق 'all' على users.import: أي إدارة، وإنشاء إدارات جديدة.
- أي نطاق آخر: إدارة المُستورِد فقط (الفارغ = إدارته)، بلا إنشاء إدارات،
  ولا إسناد دور يملك صلاحيات ليست لدى المُستورِد نفسه (منع تصعيد الصلاحيات).
"""

from __future__ import annotations

import asyncio
import io
import re
import secrets
import uuid
import zipfile
from dataclasses import dataclass, field

from email_validator import EmailNotValidError, validate_email
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.security import hash_password
from app.models.department import Department
from app.models.job_title import JobTitle
from app.models.role import Role
from app.models.user import User, UserStatus
from app.services import audit_service

SHEET_NAME = "المستخدمين"
LISTS_SHEET_NAME = "القوائم"
TEMPLATE_HEADERS = [
    "الاسم الأول",
    "الاسم الأوسط",
    "اسم العائلة",
    "البريد الإلكتروني",
    "الإدارة",
    "المسمى الوظيفي",
    "الدور في النظام",
]
MAX_ROWS = 200
MAX_FILE_BYTES = 2 * 1024 * 1024

# أسماء الأدوار كما تظهر بنموذج لاما ← أسماء الأدوار المقابلة بقاعدة البيانات
# (أول اسم موجود فعليًا يُعتمد). "ادمن" هو الاسم الحالي بالإنتاج لدور
# admin النظامي (أُعيدت تسميته من الشاشة) — نقبل الاثنين.
ROLE_ALIASES: dict[str, tuple[str, ...]] = {
    "الرئيس التنفيذي": ("executive_president",),
    "المكتب التنفيذي": ("executive_office_manager", "مدير المكتب التنفيذي"),
    "مدير إدارة (أدمن)": ("ادمن", "admin", "مسؤول إدارة"),
    "مستخدم نظام": ("مستخدم نظام",),
}
# من يُسند له دور "مدير إدارة (أدمن)" (أيًّا كان اسمه الحالي بالنظام — يُحدَّد
# عبر _resolve_role نفسها، فيبقى صحيحًا لو أُعيدت تسمية الدور) وله إدارة
# بالملف يصبح مديرها.
MANAGER_ROLE_LABEL = "مدير إدارة (أدمن)"

_NAME_MAX = 100
_DEPARTMENT_NAME_MAX = 150
_JOB_TITLE_MAX = 150
_USERNAME_MIN = 3
_USERNAME_MAX = 50
_TEMP_PASSWORD_LENGTH = 12
_UPPER = "ABCDEFGHJKLMNPQRSTUVWXYZ"
_LOWER = "abcdefghijkmnpqrstuvwxyz"
_DIGITS = "23456789"
_DEPARTMENT_CODE_RE = re.compile(r"^DEP-(\d+)$", re.IGNORECASE)


class ImportFileError(Exception):
    """خطأ بالملف ككل (ليس Excel، أعمدة خاطئة، فارغ، صفوف أكثر من الحد)."""


@dataclass
class RowError:
    row: int
    field: str | None
    message: str


class ImportValidationError(Exception):
    """صف واحد أو أكثر غير صالح — لا يُنشأ أي مستخدم."""

    def __init__(self, errors: list[RowError]):
        super().__init__("الملف يحتوي أخطاء")
        self.errors = errors


@dataclass
class ParsedRow:
    row: int
    first_name: str
    middle_name: str
    last_name: str
    email: str
    department: str
    job_title: str
    role: str


@dataclass
class CreatedUser:
    row: int
    user_id: uuid.UUID
    full_name: str
    first_name: str
    email: str
    username: str
    department: str | None
    role: str | None
    is_department_manager: bool
    # كلمة المرور المؤقتة الصريحة — لإرسالها بالبريد فقط، لا تُرجَع بالـAPI.
    temporary_password: str = field(repr=False)


@dataclass
class ImportResult:
    users: list[CreatedUser]
    departments_created: list[tuple[str, str]]  # (الاسم، الرمز)
    job_titles_created: list[str] = field(default_factory=list)


# --------------------------------------------------------------------------
# قراءة الملف
# --------------------------------------------------------------------------


def _cell_text(value: object) -> str:
    """تحويل قيمة خلية لنص مُشذَّب (الأرقام الصحيحة بدون .0)."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return " ".join(str(value).split())


def parse_workbook(content: bytes) -> list[ParsedRow]:
    """
    الهدف: قراءة ملف Excel المرفوع وتحويل صفوفه لقائمة ParsedRow.

    المدخلات: content — bytes الملف (.xlsx).
    المخرجات: list[ParsedRow] — الصفوف غير الفارغة، برقم صفها الفعلي بالملف.
    الأخطاء: ImportFileError — ملف غير صالح، شيت/أعمدة غير مطابقة للنموذج،
    ملف بلا مستخدمين، أو أكثر من MAX_ROWS صفًا.
    ملاحظة: يقرأ شيت "المستخدمين" إن وُجد وإلا الشيت الأول، ويتجاهل أي
    شيت آخر (مثل "اللجان الدائمة" — خارج نطاق هذه الميزة بقرار لاما).
    """
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError, ValueError) as exc:
        raise ImportFileError("الملف ليس ملف Excel صالحًا — استخدمي صيغة .xlsx من النموذج") from exc

    try:
        sheet = workbook[SHEET_NAME] if SHEET_NAME in workbook.sheetnames else workbook.worksheets[0]
        rows_iter = sheet.iter_rows(values_only=True)
        header = next(rows_iter, None)
        header_cells = [_cell_text(v) for v in (header or ())][: len(TEMPLATE_HEADERS)]
        if header_cells != TEMPLATE_HEADERS:
            raise ImportFileError(
                "أعمدة الملف لا تطابق النموذج. الأعمدة المطلوبة بالترتيب: " + "، ".join(TEMPLATE_HEADERS)
            )

        parsed: list[ParsedRow] = []
        for index, values in enumerate(rows_iter, start=2):
            cells = [_cell_text(v) for v in (values or ())][: len(TEMPLATE_HEADERS)]
            cells += [""] * (len(TEMPLATE_HEADERS) - len(cells))
            if not any(cells):
                continue
            if len(parsed) >= MAX_ROWS:
                raise ImportFileError(f"الحد الأقصى {MAX_ROWS} مستخدم لكل ملف — قسّمي الملف لأكثر من ملف")
            parsed.append(ParsedRow(index, *cells))
    finally:
        workbook.close()

    if not parsed:
        raise ImportFileError("الملف لا يحتوي أي مستخدمين تحت صف العناوين")
    return parsed


# --------------------------------------------------------------------------
# بيانات الدخول
# --------------------------------------------------------------------------


def username_base_from_email(email: str) -> str:
    """
    الهدف: اسم مستخدم أساسي من البريد — نفس قواعد suggestUsernameFromEmail
    بالواجهة (frontend/src/lib/credentialSuggestions.ts) حتى يكون شكل الأسماء
    موحّدًا بين الإضافة اليدوية والرفع: الجزء قبل @ بلا وسم "+"، أحرف إنجليزية
    صغيرة وأرقام و"." و"_" فقط. لو الناتج أقصر من 3 أحرف يُرجع "user".
    """
    local = email.strip().split("@")[0].split("+")[0].lower()
    base = re.sub(r"[^a-z0-9_.]", "_", local)
    base = re.sub(r"([._])[._]+", r"\1", base).strip("._")
    base = base[:_USERNAME_MAX].rstrip("._")
    return base if len(base) >= _USERNAME_MIN else "user"


def unique_username(base: str, taken: set[str]) -> str:
    """
    الهدف: أول اسم غير مأخوذ من base، base2، base3... (بحد 50 حرفًا).
    التأثيرات الجانبية: يضيف الاسم المختار إلى taken (لتفادي التكرار داخل الملف).
    """
    candidate = base
    counter = 2
    while candidate in taken:
        suffix = str(counter)
        candidate = f"{base[: _USERNAME_MAX - len(suffix)]}{suffix}"
        counter += 1
    taken.add(candidate)
    return candidate


def generate_temporary_password() -> str:
    """
    الهدف: كلمة مرور مؤقتة عشوائية آمنة (secrets) تحقق سياسة FR-UM-015:
    12 حرفًا تضم حرفًا كبيرًا وصغيرًا ورقمًا على الأقل، بلا أحرف ملتبسة
    (O/0، l/1/I) — نفس مجموعة أحرف generateTemporaryPassword بالواجهة.
    """
    alphabet = _UPPER + _LOWER + _DIGITS
    chars = [secrets.choice(_UPPER), secrets.choice(_LOWER), secrets.choice(_DIGITS)]
    chars += [secrets.choice(alphabet) for _ in range(_TEMP_PASSWORD_LENGTH - len(chars))]
    for i in range(len(chars) - 1, 0, -1):
        j = secrets.randbelow(i + 1)
        chars[i], chars[j] = chars[j], chars[i]
    return "".join(chars)


# --------------------------------------------------------------------------
# مطابقة الأدوار
# --------------------------------------------------------------------------


def _key(text: str) -> str:
    return " ".join(text.split()).casefold()


def _assignable_roles(roles: list[Role]) -> list[Role]:
    """الأدوار النظامية القابلة للإسناد عبر الملف (بدون المسؤول الأعلى وأدوار اللجان)."""
    return [r for r in roles if r.kind != "committee" and not r.is_super_admin]


def _resolve_role(label: str, roles: list[Role]) -> Role | None:
    """
    الهدف: إيجاد الدور المقصود بنص عمود "الدور في النظام".
    الترتيب: أسماء النموذج (ROLE_ALIASES) ← اسم الدور ← وصف الدور (لو
    طابق دورًا واحدًا فقط). يُرجع None لو لم يُعثر على تطابق واحد مؤكد.

    إصلاح 2026-11-01 (لاما أعادت تسمية executive_office_manager إلى "المكتب
    التنفيذي" من شاشة الأدوار): لو اسم النموذج لا يجد أسماءه البرمجية
    المتوقعة (الدور أُعيدت تسميته)، نكمل البحث بالاسم ثم الوصف بدل الرفض —
    فيعمل الرفع قبل إعادة التسمية وبعدها.
    """
    by_name = {_key(r.name): r for r in roles}

    def by_description(text: str) -> Role | None:
        matches = [r for r in roles if r.description and _key(r.description) == _key(text)]
        return matches[0] if len(matches) == 1 else None

    for alias, names in ROLE_ALIASES.items():
        if _key(alias) == _key(label):
            for name in names:
                if _key(name) in by_name:
                    return by_name[_key(name)]
            for name in names:
                if (role := by_description(name)) is not None:
                    return role
            break
    if _key(label) in by_name:
        return by_name[_key(label)]
    return by_description(label)


def role_choices(roles: list[Role]) -> list[str]:
    """
    الهدف: قائمة القيم المسموحة بعمود الدور (للنموذج ولرسائل الخطأ): أسماء
    النموذج الأربعة المتوفرة فعليًا بهذا النظام + أسماء بقية الأدوار القابلة
    للإسناد التي لا يغطيها اسم من النموذج.
    """
    assignable = _assignable_roles(roles)
    choices: list[str] = []
    covered: set[uuid.UUID] = set()
    for alias in ROLE_ALIASES:
        role = _resolve_role(alias, assignable)
        if role is not None:
            choices.append(alias)
            covered.add(role.role_id)
    choices += sorted(r.name for r in assignable if r.role_id not in covered)
    return choices


# --------------------------------------------------------------------------
# الاستيراد
# --------------------------------------------------------------------------


async def _next_department_codes(db: AsyncSession, count: int) -> list[str]:
    """رموز تلقائية DEP-001... تبدأ بعد أكبر رمز DEP-n مستخدم حاليًا (ولو لإدارة محذوفة)."""
    result = await db.execute(select(Department.code).where(Department.code.ilike("DEP-%")))
    numbers = [int(m.group(1)) for (code,) in result.all() if code and (m := _DEPARTMENT_CODE_RE.match(code))]
    start = max(numbers, default=0) + 1
    return [f"DEP-{n:03d}" for n in range(start, start + count)]


async def import_users(db: AsyncSession, *, actor: User, rows: list[ParsedRow]) -> ImportResult:
    """
    الهدف:
    التحقق من كل صفوف الملف ثم إنشاء المستخدمين (والإدارات الجديدة) بمعاملة
    واحدة — الكل أو لا شيء.

    المدخلات:
    - db: جلسة قاعدة البيانات.
    - actor: المستخدم المُستورِد (يملك users.import — مفحوص بطبقة الـAPI).
    - rows: ناتج parse_workbook.

    المخرجات: ImportResult — المستخدمون المنشأون (مع كلمات المرور المؤقتة
    للإرسال بالبريد فقط) والإدارات المنشأة.

    التأثيرات الجانبية: يكتب users/departments/audit_logs ويعمل commit.
    كل مستخدم must_change_password=True (FR-UM-016).

    الأخطاء: ImportValidationError بكل أخطاء الصفوف (لا يُكتب شيء).
    """
    full_scope = actor.scope_for("users.import") == "all"
    actor_permissions = actor.permission_codes
    if not full_scope and actor.dep_id is None:
        raise ImportFileError("حسابك غير مرتبط بإدارة — لا يمكنك رفع مستخدمين إلا لإدارتك")

    emails_in_file = {r.email.casefold() for r in rows if r.email}
    roles = _assignable_roles(list((await db.execute(select(Role))).scalars().all()))
    manager_role = _resolve_role(MANAGER_ROLE_LABEL, roles)
    job_titles = {_key(j.name): j for j in (await db.execute(select(JobTitle))).scalars().all()}
    departments = {
        _key(d.name): d
        for d in (await db.execute(select(Department).where(Department.deleted_at.is_(None)))).scalars().all()
    }
    existing_emails = set(
        (
            await db.execute(
                select(func.lower(User.email)).where(
                    User.deleted_at.is_(None), func.lower(User.email).in_(emails_in_file or {""})
                )
            )
        ).scalars()
    )
    taken_usernames = set(
        (await db.execute(select(func.lower(User.username)).where(User.deleted_at.is_(None)))).scalars()
    )

    errors: list[RowError] = []
    seen_emails: dict[str, int] = {}
    # (الصف، البريد، الدور، مفتاح المسمى، مفتاح الإدارة) — المسمى والإدارة
    # بالمفتاح لأن الجديد منهما لا يُنشأ إلا بعد نجاح التحقق من كل الصفوف.
    resolved: list[tuple[ParsedRow, str, Role | None, str | None, str | None]] = []
    can_create_job_titles = full_scope or "job_titles.create" in actor_permissions
    new_job_titles: dict[str, str] = {}  # key → الاسم كما كُتب أول مرة
    new_departments: dict[str, str] = {}  # key → الاسم كما كُتب أول مرة
    managers: dict[str, int] = {}  # key الإدارة → رقم صف مديرها بالملف
    department_first_row: dict[str, int] = {}

    for row in rows:
        row_errors: list[RowError] = []

        for value, field_name in (
            (row.first_name, "الاسم الأول"),
            (row.middle_name, "الاسم الأوسط"),
            (row.last_name, "اسم العائلة"),
        ):
            if not value:
                row_errors.append(RowError(row.row, field_name, f"{field_name} مطلوب"))
            elif len(value) > _NAME_MAX:
                row_errors.append(RowError(row.row, field_name, f"{field_name} أطول من {_NAME_MAX} حرف"))

        email = ""
        if not row.email:
            row_errors.append(RowError(row.row, "البريد الإلكتروني", "البريد الإلكتروني مطلوب"))
        else:
            try:
                email = validate_email(row.email, check_deliverability=False).normalized
            except EmailNotValidError:
                row_errors.append(RowError(row.row, "البريد الإلكتروني", f"صيغة البريد غير صحيحة: {row.email}"))
            else:
                email_key = email.casefold()
                if email_key in seen_emails:
                    row_errors.append(
                        RowError(row.row, "البريد الإلكتروني", f"البريد مكرر بالملف (نفس الصف {seen_emails[email_key]})")
                    )
                elif email_key in existing_emails:
                    row_errors.append(RowError(row.row, "البريد الإلكتروني", "البريد مسجّل مسبقًا لمستخدم آخر بالنظام"))
                seen_emails.setdefault(email_key, row.row)

        role: Role | None = None
        if row.role:
            role = _resolve_role(row.role, roles)
            if role is None:
                row_errors.append(
                    RowError(
                        row.row,
                        "الدور في النظام",
                        f"الدور «{row.role}» غير موجود. القيم المسموحة: " + "، ".join(role_choices(roles)),
                    )
                )
            elif not full_scope and not role.permission_codes <= actor_permissions:
                row_errors.append(
                    RowError(row.row, "الدور في النظام", f"لا تملك صلاحية إسناد الدور «{row.role}»")
                )

        job_title_key: str | None = None
        if row.job_title:
            job_title_key = _key(row.job_title)
            if len(row.job_title) > _JOB_TITLE_MAX:
                row_errors.append(
                    RowError(row.row, "المسمى الوظيفي", f"المسمى الوظيفي أطول من {_JOB_TITLE_MAX} حرف")
                )
            elif job_title_key not in job_titles:
                if can_create_job_titles:
                    new_job_titles.setdefault(job_title_key, row.job_title)
                else:
                    row_errors.append(
                        RowError(
                            row.row,
                            "المسمى الوظيفي",
                            f"المسمى «{row.job_title}» غير موجود، ولا تملك صلاحية إضافة مسميات وظيفية",
                        )
                    )

        department_key: str | None = None
        if row.department:
            department_key = _key(row.department)
            if len(row.department) > _DEPARTMENT_NAME_MAX:
                row_errors.append(RowError(row.row, "الإدارة", f"اسم الإدارة أطول من {_DEPARTMENT_NAME_MAX} حرف"))
            elif department_key not in departments:
                if full_scope:
                    new_departments.setdefault(department_key, row.department)
                else:
                    row_errors.append(
                        RowError(row.row, "الإدارة", f"الإدارة «{row.department}» غير موجودة، ولا تملك صلاحية إنشاء إدارات")
                    )
        elif not full_scope and actor.department is not None:
            department_key = _key(actor.department.name)

        if not full_scope and department_key is not None and not any(e.field == "الإدارة" for e in row_errors):
            target = departments.get(department_key)
            if target is None or target.dep_id != actor.dep_id:
                row_errors.append(RowError(row.row, "الإدارة", "يمكنك إضافة مستخدمين لإدارتك فقط"))

        if department_key is not None:
            department_first_row.setdefault(department_key, row.row)
            if role is not None and manager_role is not None and role.role_id == manager_role.role_id:
                existing = departments.get(department_key)
                if department_key in managers:
                    row_errors.append(
                        RowError(row.row, "الدور في النظام", f"للإدارة مدير آخر بالملف (الصف {managers[department_key]})")
                    )
                elif existing is not None and existing.manager_user_id is not None:
                    row_errors.append(
                        RowError(
                            row.row,
                            "الدور في النظام",
                            f"الإدارة «{existing.name}» لها مدير — تغيير المدير يتم من صفحة الإدارات",
                        )
                    )
                else:
                    managers[department_key] = row.row

        errors += row_errors
        if not row_errors:
            resolved.append((row, email, role, job_title_key, department_key))

    for key, name in new_departments.items():
        if key not in managers:
            errors.append(
                RowError(
                    department_first_row[key],
                    "الإدارة",
                    f"الإدارة «{name}» جديدة وستُنشأ تلقائيًا، وتحتاج مديرًا: "
                    f"أضيفي صفًا لها دوره «{MANAGER_ROLE_LABEL}»",
                )
            )

    if errors:
        errors.sort(key=lambda e: e.row)
        raise ImportValidationError(errors)

    # ---- الإنشاء (معاملة واحدة) ----
    created_job_titles: list[str] = []
    for key, name in new_job_titles.items():
        job_title = JobTitle(name=name)
        db.add(job_title)
        job_titles[key] = job_title
        created_job_titles.append(name)
    if new_job_titles:
        await db.flush()
        for name in created_job_titles:
            await audit_service.log_action(
                db,
                actor_user_id=actor.user_id,
                action_type="create",
                target_type="job_title",
                target_id=job_titles[_key(name)].job_title_id,
                metadata={"name": name, "source": "excel_import"},
            )

    codes = await _next_department_codes(db, len(new_departments))
    created_departments: list[tuple[str, str]] = []
    for (key, name), code in zip(new_departments.items(), codes):
        department = Department(name=name, code=code, description=None)
        db.add(department)
        departments[key] = department
        created_departments.append((name, code))
    if new_departments:
        await db.flush()

    passwords = [generate_temporary_password() for _ in resolved]
    # bcrypt بطيء عمدًا — التجزئة بالتوازي على خيوط منفصلة بدل التسلسل.
    hashes = await asyncio.gather(*(asyncio.to_thread(hash_password, p) for p in passwords))

    created: list[CreatedUser] = []
    for (row, email, role, job_title_key, department_key), password, password_hash in zip(resolved, passwords, hashes):
        department = departments.get(department_key) if department_key else None
        job_title = job_titles.get(job_title_key) if job_title_key else None
        username = unique_username(username_base_from_email(email), taken_usernames)
        user = User(
            first_name=row.first_name,
            middle_name=row.middle_name,
            last_name=row.last_name,
            username=username,
            email=email,
            password_hash=password_hash,
            role_id=role.role_id if role else None,
            dep_id=department.dep_id if department else None,
            job_title_id=job_title.job_title_id if job_title else None,
            status=UserStatus.active,
            must_change_password=True,
        )
        db.add(user)
        await db.flush()

        is_manager = bool(department_key and managers.get(department_key) == row.row)
        if is_manager and department is not None:
            department.manager_user_id = user.user_id

        await audit_service.log_action(
            db,
            actor_user_id=actor.user_id,
            action_type="create",
            target_type="user",
            target_id=user.user_id,
            metadata={"username": username, "role": role.name if role else None, "source": "excel_import"},
        )
        created.append(
            CreatedUser(
                row=row.row,
                user_id=user.user_id,
                full_name=f"{row.first_name} {row.middle_name} {row.last_name}",
                first_name=row.first_name,
                email=email,
                username=username,
                department=department.name if department else None,
                role=row.role or None,
                is_department_manager=is_manager,
                temporary_password=password,
            )
        )

    for name, code in created_departments:
        department = departments[_key(name)]
        await audit_service.log_action(
            db,
            actor_user_id=actor.user_id,
            action_type="create",
            target_type="department",
            target_id=department.dep_id,
            metadata={
                "name": name,
                "code": code,
                "manager_user_id": str(department.manager_user_id),
                "source": "excel_import",
            },
        )

    await db.commit()
    return ImportResult(
        users=created, departments_created=created_departments, job_titles_created=created_job_titles
    )


# --------------------------------------------------------------------------
# النموذج
# --------------------------------------------------------------------------


async def build_template(db: AsyncSession) -> bytes:
    """
    الهدف: نموذج Excel فارغ بنفس الأعمدة المطلوبة، مع قوائم منسدلة محدثة من
    النظام: الأدوار (role_choices) والمسميات الوظيفية، وشيت "القوائم" يعرض
    الإدارات الحالية للمرجع (عمود الإدارة يقبل أيضًا اسم إدارة جديدة).
    المخرجات: bytes ملف .xlsx.
    """
    roles = _assignable_roles(list((await db.execute(select(Role))).scalars().all()))
    titles = sorted((await db.execute(select(JobTitle.name))).scalars().all())
    department_names = sorted(
        (await db.execute(select(Department.name).where(Department.deleted_at.is_(None)))).scalars().all()
    )
    choices = role_choices(roles)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    sheet.sheet_view.rightToLeft = True
    sheet.append(TEMPLATE_HEADERS)
    header_fill = PatternFill("solid", fgColor="105EA7")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for column, width in zip("ABCDEFG", (16, 16, 16, 30, 24, 24, 22)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"

    lists = workbook.create_sheet(LISTS_SHEET_NAME)
    lists.sheet_view.rightToLeft = True
    lists.append(["الأدوار", "المسميات الوظيفية", "الإدارات الحالية"])
    for cell in lists[1]:
        cell.font = Font(bold=True)
    for i in range(max(len(choices), len(titles), len(department_names))):
        lists.append(
            [
                choices[i] if i < len(choices) else None,
                titles[i] if i < len(titles) else None,
                department_names[i] if i < len(department_names) else None,
            ]
        )
    for column in "ABC":
        lists.column_dimensions[column].width = 28

    last_data_row = MAX_ROWS + 1
    if choices:
        role_validation = DataValidation(
            type="list", formula1=f"'{LISTS_SHEET_NAME}'!$A$2:$A${len(choices) + 1}", allow_blank=True
        )
        role_validation.add(f"G2:G{last_data_row}")
        sheet.add_data_validation(role_validation)
    if titles:
        title_validation = DataValidation(
            type="list", formula1=f"'{LISTS_SHEET_NAME}'!$B$2:$B${len(titles) + 1}", allow_blank=True
        )
        title_validation.add(f"F2:F{last_data_row}")
        sheet.add_data_validation(title_validation)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
